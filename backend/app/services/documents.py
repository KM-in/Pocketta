from __future__ import annotations
import asyncio
from contextlib import closing
import hashlib
import json
from pathlib import Path
from threading import Lock
from ..models import Transcript, TranscriptSegment
from .gateway import ModelGateway
from .scheduler import CourseScheduler
from .processes import ProcessRegistry

_PDFIUM_LOCK = Lock()  # PDFium native calls must not overlap across document jobs.

VISION_PROMPT = "Transcribe visible text and describe only clearly supported diagrams and tables. Do not infer missing facts. Return transcription, description, and uncertain (true if anything material is unclear)."


def needs_vision(
    text: str, visual_coverage: float, structured_visual: bool = False
) -> bool:
    return len(text.strip()) < 80 or visual_coverage >= 0.15 or structured_visual


def render_pdf_page(path: Path, page_index: int, output: Path):
    import pypdfium2 as pdfium

    cache = output.with_suffix(".json")
    if output.is_file() and cache.is_file():
        try:
            return tuple(json.loads(cache.read_text()))
        except (ValueError, OSError):
            pass  # A partial cache from a crash can be rebuilt from the original.
    try:
        with _PDFIUM_LOCK, closing(pdfium.PdfDocument(path)) as pdf:
            page = pdf[page_index]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_bounded()
                finally:
                    textpage.close()
                area = max(1, page.get_width() * page.get_height())
                coverage = 0.0
                paths = 0
                for obj in page.get_objects():
                    if obj.type == pdfium.raw.FPDF_PAGEOBJ_IMAGE:
                        left, bottom, right, top = obj.get_bounds()
                        coverage += max(0, right - left) * max(0, top - bottom) / area
                    elif obj.type == pdfium.raw.FPDF_PAGEOBJ_PATH:
                        paths += 1
                bitmap = page.render(
                    scale=min(2, 1800 / max(page.get_width(), page.get_height()))
                )
                try:
                    bitmap.to_pil().save(output)
                finally:
                    bitmap.close()
                result = (text, coverage, paths >= 12)
                cache.write_text(json.dumps(result))
                return result
            finally:
                page.close()
    except Exception as error:
        raise ValueError(
            "Cannot read PDF page; check whether the file is corrupt or password protected"
        ) from error


def pdf_length(path: Path):
    import pypdfium2 as pdfium

    try:
        with _PDFIUM_LOCK, closing(pdfium.PdfDocument(path)) as pdf:
            return len(pdf)
    except Exception as error:
        raise ValueError(
            "Cannot open PDF; encrypted and corrupt PDFs are unsupported"
        ) from error


def pptx_text(path: Path):
    from pptx import Presentation

    result = []
    for slide in Presentation(path).slides:
        parts, visual = [], False
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text)
            if shape.has_table:
                visual = True
                parts.extend(
                    " | ".join(cell.text for cell in row.cells)
                    for row in shape.table.rows
                )
            if shape.has_chart or shape.shape_type in (3, 6, 13):
                visual = True
        result.append(("\n".join(parts), visual))
    return result


class DocumentProcessor:
    def __init__(
        self, settings, store, scheduler=None, coordinator=None, knowledge=None
    ):
        self.settings, self.store = settings, store
        self.scheduler = scheduler or CourseScheduler(settings)
        self.owns_scheduler = scheduler is None
        self.coordinator = coordinator
        if knowledge is None:
            from .knowledge import KnowledgeService

            knowledge = KnowledgeService(settings, store, coordinator=coordinator)
        self.knowledge = knowledge
        self.queue = asyncio.Queue()
        self.pending = set()
        self.active = {}
        self.cancelled = set()
        self.processes = ProcessRegistry()

    async def blocking(self, fn, *args):
        # Await native work on cancellation before deleting its files.
        task = asyncio.create_task(asyncio.to_thread(fn, *args))
        try:
            return await asyncio.shield(task)
        except asyncio.CancelledError:
            try:
                await task
            except Exception:
                pass
            raise

    async def start(self):
        with self.store.db._connect() as c:
            ids = [
                (r["id"], r["work_kind"])
                for r in c.execute(
                    "SELECT id,work_kind FROM sources WHERE type!='recording' AND status IN ('queued','extracting','indexing','generating')"
                )
            ]
        for id, kind in ids:
            self.enqueue(id, generation_only=kind == "generate")
        if self.owns_scheduler:
            await self.scheduler.start()

    async def stop(self):
        if self.owns_scheduler:
            self.scheduler.running = False
        for id in list(self.pending):
            await asyncio.to_thread(self.processes.terminate, id)
            await self.scheduler.cancel(id)
            if self.store.source(id):
                self.store.update(id, "queued")
        if self.owns_scheduler:
            await self.scheduler.stop()

    def enqueue(self, id, generation_only=False):
        if (
            id in self.pending
            or id in self.active
            or f"knowledge:{id}" in self.scheduler.jobs
        ):
            raise ValueError("Source is already processing")
        self.cancelled.discard(id)
        self.store.update(id, "queued")
        with self.store.db._connect() as c:
            c.execute(
                "UPDATE sources SET work_kind=? WHERE id=?",
                ("generate" if generation_only else "extract", id),
            )
        self.pending.add(id)
        self.queue.put_nowait((id, generation_only))
        course_id = self.store.source(id)["course_id"]
        self.scheduler.submit(
            id,
            course_id,
            lambda: self._run_job(id, generation_only),
            lambda: self._job_done(id),
        )

    async def cancel(self, id):
        self.cancelled.add(id)
        self.store.update(id, "cancelled")
        await asyncio.to_thread(self.processes.terminate, id)
        await self.scheduler.cancel(id)

    def check(self, id):
        if id in self.cancelled or not self.store.source(id):
            raise asyncio.CancelledError()

    def _job_done(self, id):
        self.pending.discard(id)
        self.cancelled.discard(id)
        self.active.pop(id, None)
        self.queue.get_nowait()
        self.queue.task_done()

    async def _run_job(self, id, generation_only):
        self.active[id] = asyncio.current_task()
        try:
            self.check(id)
            await self.process(id, generation_only)
        except asyncio.CancelledError:
            raise
        except Exception as error:
            self.store.update(id, "failed", error=str(error)[-1000:])

    async def process(self, id, generation_only):
        source = self.store.source(id)
        path = Path(self.store.path(id))
        gateway = ModelGateway(
            self.settings, self.store.db, id, coordinator=self.coordinator
        )
        if not generation_only:
            self.store.update(id, "extracting", 5)
            slides = None
            pdf_path = path
            if source["type"] == "pptx":
                if not self.settings.executable_available(
                    self.settings.libreoffice_path
                ):
                    raise ValueError(
                        "Install LibreOffice and set LIBREOFFICE_PATH to soffice to render PPTX slides"
                    )
                slides = await self.blocking(pptx_text, path)
                self.check(id)
                pdf_path = path.with_suffix(".pdf")
                if not pdf_path.exists():
                    await self.blocking(
                        self.processes.run,
                        id,
                        [
                            self.settings.libreoffice_path,
                            f'-env:UserInstallation={path.parent.joinpath("lo-profile").as_uri()}',
                            "--headless",
                            "--convert-to",
                            "pdf",
                            "--outdir",
                            str(path.parent),
                            str(path),
                        ],
                        None,
                        120,
                    )
                self.check(id)
            count = await self.blocking(pdf_length, pdf_path)
            if count > self.settings.max_document_pages:
                raise ValueError(
                    f"Document exceeds {self.settings.max_document_pages} pages"
                )
            if slides is not None and len(slides) != count:
                raise ValueError(
                    "Rendered slide count does not match PPTX; cannot guarantee citation locations"
                )
            for index in range(count):
                self.check(id)
                image_path = path.parent / f"page-{index+1}.png"
                native, coverage, structured = await self.blocking(
                    render_pdf_page, pdf_path, index, image_path
                )
                self.check(id)
                if slides is not None:
                    native, slide_visual = slides[index]
                    structured = structured or slide_visual
                location = f'{"slide" if slides is not None else "page"}:{index+1}'
                unit_id = self.store.save_unit(
                    id, location, native, native, str(image_path)
                )
                if needs_vision(native, coverage, structured):
                    cache_key = hashlib.sha256(
                        (
                            str(source["version"])
                            + self.settings.cerebras_vision_model
                            + VISION_PROMPT
                        ).encode()
                        + image_path.read_bytes()
                    ).hexdigest()
                    unit = next(u for u in self.store.units(id) if u["id"] == unit_id)
                    cached = (
                        json.loads(unit["vision_json"]) if unit["vision_json"] else None
                    )
                    if cached and cached.get("cache_key") == cache_key:
                        vision = cached
                    else:
                        result = await gateway.analyse_image(image_path, VISION_PROMPT)
                        vision = {**result.model_dump(), "cache_key": cache_key}
                    self.check(id)
                    text = "\n".join(
                        p
                        for p in (
                            native,
                            vision["transcription"] if len(native.strip()) < 80 else "",
                            vision["description"],
                        )
                        if p
                    )
                    with self.store.db._connect() as c:
                        c.execute(
                            "UPDATE content_units SET text=?,vision_json=?,uncertain=? WHERE id=?",
                            (text, json.dumps(vision), vision["uncertain"], unit_id),
                        )
                self.store.update(id, "extracting", int(10 + 70 * (index + 1) / count))
            self.store.update(id, "indexing", 85)
        self.check(id)
        await self.knowledge.index_source(id)
        self.check(id)
        self.store.update(id, "ready", 100)
        # Extraction stays usable even when generation is unavailable. Generation is explicit.
        if generation_only:
            units = self.store.units(id)
            reliable = [u for u in units if u["text"].strip() and not u["uncertain"]]
            if not reliable:
                raise ValueError(
                    "No reliable content available; inspect extraction warnings"
                )
            self.store.update(id, "generating", 90)
            transcript = Transcript(
                duration_ms=0,
                segments=[
                    TranscriptSegment(
                        id=u["id"],
                        start_ms=0,
                        end_ms=0,
                        text=u["text"],
                        confidence=1,
                        uncertain=False,
                    )
                    for u in reliable
                ],
            )
            pack = await gateway.generate(transcript)
            self.check(id)
            pack.validate_evidence(transcript)
            data = pack.model_dump()
            for key in ("notes", "concepts", "flashcards", "quiz"):
                for item in data[key]:
                    item["content_unit_ids"] = item.pop("segment_ids")
            with self.store.db._connect() as c:
                c.execute(
                    "UPDATE sources SET pack_json=? WHERE id=?", (json.dumps(data), id)
                )
            self.store.update(id, "ready", 100)
