from pathlib import Path
from unittest.mock import AsyncMock
import json
import uuid
import pytest
from fastapi.testclient import TestClient
from backend.app.config import Settings
from backend.app.courses import CourseStore, IMPORT_COURSE
from backend.app.database import Database
from backend.app.main import create_app
from backend.app.services.documents import (
    DocumentProcessor,
    needs_vision,
    render_pdf_page,
    pptx_text,
)
from backend.app.services.gateway import ModelGateway, VisionResult
from backend.tests.test_models import transcript, study_pack


def pdf_bytes(text=""):
    content = f"BT /F1 12 Tf 30 750 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(content)).encode()
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
    ]
    output, offsets = b"%PDF-1.4\n", [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    start = len(output)
    output += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(
        f"{o:010d} 00000 n \n".encode() for o in offsets[1:]
    )
    return (
        output
        + f"trailer\n<< /Root 1 0 R /Size 6 >>\nstartxref\n{start}\n%%EOF".encode()
    )


@pytest.fixture
def store(tmp_path):
    db = Database(tmp_path / "db.sqlite")
    db.initialize()
    result = CourseStore(db)
    result.initialize()
    return result


def add_document(store, tmp_path, content=None, kind="pdf"):
    course = store.create_course("Operating systems", "")
    id = str(uuid.uuid4())
    directory = tmp_path / id
    directory.mkdir()
    path = directory / f"source.{kind}"
    path.write_bytes(
        content if content is not None else pdf_bytes("Processes share resources. " * 8)
    )
    with store.db._connect() as c:
        c.execute(
            "INSERT INTO sources(id,course_id,name,type,file_path,status,created_at) VALUES (?,?,?,?,?,'queued','now')",
            (id, course["id"], "Textbook", kind, str(path)),
        )
    return id, path


def test_migration_preserves_legacy_and_stable_ids(tmp_path):
    db = Database(tmp_path / "db.sqlite")
    db.initialize()
    db.create_lecture("legacy-id", "old.mp3", tmp_path / "old.mp3")
    db.save_transcript("legacy-id", transcript())
    db.save_study_pack("legacy-id", study_pack())
    store = CourseStore(db)
    store.initialize()
    before = store.source("legacy-id")
    store.initialize()
    after = store.source("legacy-id")
    assert len(store.sources(IMPORT_COURSE)) == 1
    assert before["content_units"] == after["content_units"]
    assert after["study_pack"]["notes"][0]["content_unit_ids"] == [
        after["content_units"][0]["id"]
    ]
    assert db.get_lecture("legacy-id").study_pack == study_pack()
    revised = transcript()
    revised.segments[0].text = "Corrected evidence"
    db.replace_transcript_and_invalidate_pack("legacy-id", revised)
    corrected = store.source("legacy-id")
    assert corrected["content_units"][0]["id"] == after["content_units"][0]["id"]
    assert corrected["content_units"][0]["text"] == "Corrected evidence"
    assert corrected["study_pack"] is None


def test_course_crud_isolation_and_upload_cleanup(tmp_path, monkeypatch):
    app = create_app(Settings(_env_file=None, pocketta_data_dir=tmp_path))
    monkeypatch.setattr(app.state.documents, "enqueue", lambda *a, **k: None)
    with TestClient(app) as client:
        assert client.post("/api/courses", json={"title": "   "}).status_code == 422
        a = client.post("/api/courses", json={"title": " A "}).json()
        b = client.post("/api/courses", json={"title": "B"}).json()
        assert a["title"] == "A"
        assert (
            client.post(
                "/api/courses/missing/sources", files={"file": ("a.pdf", b"bad")}
            ).status_code
            == 404
        )
        assert (
            client.post(
                f'/api/courses/{a["id"]}/sources', files={"file": ("empty.pdf", b"")}
            ).status_code
            == 400
        )
        assert list((tmp_path / "sources").iterdir()) == []
        source = client.post(
            f'/api/courses/{a["id"]}/sources', files={"file": ("a.pdf", pdf_bytes())}
        ).json()
        assert client.get(f'/api/courses/{b["id"]}/sources').json() == []
        assert client.delete(f'/api/courses/{a["id"]}').status_code == 409
        assert "file_path" not in source
        assert client.delete(f'/api/sources/{source["id"]}').status_code == 204
        assert client.get(f'/api/sources/{source["id"]}/asset').status_code == 404
        assert client.delete(f'/api/courses/{a["id"]}').status_code == 204


@pytest.mark.asyncio
async def test_native_pdf_skips_vision_and_retains_original(
    store, tmp_path, monkeypatch
):
    id, path = add_document(store, tmp_path)
    vision = AsyncMock()
    monkeypatch.setattr(ModelGateway, "analyse_image", vision)
    worker = DocumentProcessor(
        Settings(_env_file=None, pocketta_data_dir=tmp_path), store
    )
    await worker.process(id, False)
    vision.assert_not_called()
    source = store.source(id)
    assert source["status"] == "ready"
    assert source["content_units"][0]["location"] == "page:1"
    assert "Processes" in source["content_units"][0]["text"]
    assert Path(source["content_units"][0]["image_path"]).is_file()
    assert path.is_file()
    assert needs_vision("text " * 30, 0.2)
    assert needs_vision("text " * 30, 0, True)


@pytest.mark.asyncio
async def test_scan_vision_cache_and_stable_ids(store, tmp_path, monkeypatch):
    id, path = add_document(store, tmp_path, pdf_bytes())
    vision = AsyncMock(
        return_value=VisionResult(
            transcription="Scanned evidence",
            description="A process diagram",
            uncertain=False,
        )
    )
    monkeypatch.setattr(ModelGateway, "analyse_image", vision)
    worker = DocumentProcessor(
        Settings(_env_file=None, pocketta_data_dir=tmp_path, model_provider="cerebras"),
        store,
    )
    await worker.process(id, False)
    before = store.units(id)
    await worker.process(id, False)
    assert vision.await_count == 1
    assert before == store.units(id)
    assert "Scanned evidence" in before[0]["text"]
    assert before[0]["native_text"] == ""
    assert json.loads(before[0]["vision_json"])["cache_key"]


@pytest.mark.asyncio
async def test_generation_rejects_foreign_and_invented_ids(
    store, tmp_path, monkeypatch
):
    id, path = add_document(store, tmp_path)
    await DocumentProcessor(
        Settings(_env_file=None, pocketta_data_dir=tmp_path), store
    ).process(id, False)
    unit = store.units(id)[0]
    grounded = transcript()
    grounded.segments[0].id = unit["id"]
    gateway = ModelGateway(
        Settings(_env_file=None, model_provider="cerebras", cerebras_api_key="test"),
        store.db,
        id,
    )
    complete = AsyncMock(return_value=study_pack(str(uuid.uuid4())).model_dump_json())
    monkeypatch.setattr(gateway, "_complete", complete)
    with pytest.raises(RuntimeError, match="Unknown transcript segment IDs"):
        await gateway.generate(grounded)
    assert complete.await_count == 2
    complete.return_value = study_pack(unit["id"]).model_dump_json()
    assert (await gateway.generate(grounded)).notes[0].segment_ids == [unit["id"]]


@pytest.mark.asyncio
async def test_missing_tools_corrupt_files_and_recovery(store, tmp_path):
    id, path = add_document(store, tmp_path, b"corrupt")
    worker = DocumentProcessor(
        Settings(
            _env_file=None,
            pocketta_data_dir=tmp_path,
            libreoffice_path="missing-soffice",
        ),
        store,
    )
    await worker.start()
    await worker.queue.join()
    assert store.source(id)["status"] == "failed"
    assert "encrypted" in store.source(id)["error"]
    await worker.stop()
    ppt_id, _ = add_document(store, tmp_path, b"invalid pptx", "pptx")
    with pytest.raises(ValueError, match="LibreOffice"):
        await worker.process(ppt_id, False)


def test_pptx_text_preserves_slide_boundaries(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    for title in ("Processes", "Scheduling"):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
    path = tmp_path / "slides.pptx"
    deck.save(path)
    assert [text for text, _ in pptx_text(path)] == ["Processes\n", "Scheduling\n"]


@pytest.mark.asyncio
async def test_generation_recovery_and_citation_resolution(tmp_path, monkeypatch):
    app = create_app(Settings(_env_file=None, pocketta_data_dir=tmp_path))
    app.state.settings.prepare_directories()
    app.state.database.initialize()
    app.state.courses.initialize()
    store = app.state.courses
    id, path = add_document(store, tmp_path)
    unit_id = store.save_unit(
        id, "page:7", "Known page evidence", "Known page evidence"
    )
    with store.db._connect() as c:
        c.execute(
            "UPDATE sources SET status='generating',work_kind='generate' WHERE id=?",
            (id,),
        )
    generate = AsyncMock(return_value=study_pack(unit_id))
    monkeypatch.setattr(ModelGateway, "generate", generate)
    # Startup must resume generation rather than replacing it with an extraction-only job.
    await app.state.documents.start()
    await app.state.documents.queue.join()
    await app.state.documents.stop()
    assert store.source(id)["study_pack"]["notes"][0]["content_unit_ids"] == [unit_id]
    app = create_app(Settings(_env_file=None, pocketta_data_dir=tmp_path))
    with TestClient(app) as client:
        citation = client.get(f"/api/citations/{unit_id}").json()
        assert citation["page_no"] == 7
        assert citation["source_id"] == id
        exported = client.get(f"/api/sources/{id}/export.md")
        assert (
            exported.status_code == 200
            and "page:7" in exported.text
            and unit_id in exported.text
        )
        assert (
            client.get(f"/api/sources/{uuid.uuid4()}/units/{unit_id}/image").status_code
            == 404
        )
        assert client.get("/api/citations/missing").status_code == 404


@pytest.mark.asyncio
async def test_cancellation_waits_for_native_work_before_cleanup(
    store, tmp_path, monkeypatch
):
    import asyncio
    import threading
    from backend.app.services import documents

    id, path = add_document(store, tmp_path)
    started, release = threading.Event(), threading.Event()
    original = documents.render_pdf_page

    def slow_render(*args):
        started.set()
        assert release.wait(timeout=5)
        return original(*args)

    monkeypatch.setattr(documents, "render_pdf_page", slow_render)
    worker = DocumentProcessor(
        Settings(_env_file=None, pocketta_data_dir=tmp_path), store
    )
    await worker.start()
    assert await asyncio.to_thread(started.wait, 5)
    cancellation = asyncio.create_task(worker.cancel(id))
    await asyncio.sleep(0.01)
    assert not cancellation.done()
    release.set()
    await cancellation
    await worker.queue.join()
    assert store.source(id)["status"] == "cancelled"
    assert store.units(id) == []
    assert path.exists()
    await worker.stop()


@pytest.mark.asyncio
async def test_gateway_retries_rate_limits_records_usage_and_hides_errors(
    store, tmp_path, monkeypatch
):
    import httpx
    from backend.app.services import gateway as module

    id, _ = add_document(store, tmp_path)
    calls = []

    def respond(request):
        calls.append(request)
        if len(calls) == 1:
            return httpx.Response(429)
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "message": {
                            "content": json.dumps(
                                {
                                    "transcription": "Text",
                                    "description": "Diagram",
                                    "uncertain": False,
                                }
                            )
                        }
                    }
                ],
                "usage": {"prompt_tokens": 12, "completion_tokens": 8},
            },
        )

    original = httpx.AsyncClient
    monkeypatch.setattr(
        module.httpx,
        "AsyncClient",
        lambda **kwargs: original(transport=httpx.MockTransport(respond), **kwargs),
    )
    monkeypatch.setattr(module.asyncio, "sleep", AsyncMock())
    image = tmp_path / "image.png"
    image.write_bytes(b"png")
    gateway = ModelGateway(
        Settings(
            _env_file=None, model_provider="cerebras", cerebras_api_key="secret-value"
        ),
        store.db,
        id,
    )
    result = await gateway.analyse_image(image, "Read only visible evidence")
    assert result.transcription == "Text" and len(calls) == 2
    payload = json.loads(calls[-1].content)
    assert payload["messages"][0]["content"][1]["image_url"]["url"].startswith(
        "data:image/png;base64,"
    )
    assert (
        payload["response_format"]["json_schema"]["schema"]["additionalProperties"]
        is False
    )
    with store.db._connect() as c:
        row = dict(c.execute("SELECT * FROM model_runs").fetchone())
    assert (
        row["status"] == "completed"
        and json.loads(row["usage_json"])["prompt_tokens"] == 12
    )
    assert "secret-value" not in json.dumps(row)


def test_encrypted_pdf_fails_with_actionable_error(tmp_path):
    import hashlib
    import struct
    from backend.app.services.documents import pdf_length

    # A genuine PDF Standard Security (revision 2) fixture; no test-only PDF dependency.
    def rc4(key, data):
        state = list(range(256))
        j = 0
        for i in range(256):
            j = (j + state[i] + key[i % len(key)]) % 256
            state[i], state[j] = state[j], state[i]
        i = j = 0
        output = bytearray()
        for byte in data:
            i = (i + 1) % 256
            j = (j + state[i]) % 256
            state[i], state[j] = state[j], state[i]
            output.append(byte ^ state[(state[i] + state[j]) % 256])
        return bytes(output)

    padding = bytes.fromhex(
        "28bf4e5e4e758a4164004e56fffa01082e2e00b6d0683e802f0ca9fe6453697a"
    )
    pad = lambda value: (value + padding)[:32]
    file_id = hashlib.md5(b"pocketta-test").digest()
    owner = rc4(hashlib.md5(pad(b"owner")).digest()[:5], pad(b"password"))
    key = hashlib.md5(
        pad(b"password") + owner + struct.pack("<i", -4) + file_id
    ).digest()[:5]
    user = rc4(key, padding)
    raw = pdf_bytes("Secret chapter")
    encrypt = f"6 0 obj\n<< /Filter /Standard /V 1 /R 2 /Length 40 /O <{owner.hex()}> /U <{user.hex()}> /P -4 >>\nendobj\n".encode()
    # PDFium reconstructs xrefs for this small fixture, preserving the encryption dictionary.
    raw = raw.replace(b"xref\n0 6", encrypt + b"xref\n0 6", 1)
    raw = raw.replace(
        b"/Root 1 0 R",
        f"/Encrypt 6 0 R /ID [<{file_id.hex()}> <{file_id.hex()}>] /Root 1 0 R".encode(),
    )
    path = tmp_path / "encrypted.pdf"
    path.write_bytes(raw)
    with pytest.raises(ValueError, match="encrypted and corrupt"):
        pdf_length(path)
