import shutil
import subprocess
import uuid
from datetime import UTC, datetime

import pytest
from pptx import Presentation

from backend.app.config import Settings
from backend.app.courses import CourseStore
from backend.app.database import Database
from backend.app.services.documents import DocumentProcessor
from backend.app.services.knowledge import KnowledgeService


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice is required")
@pytest.mark.asyncio
async def test_legacy_ppt_converts_and_keeps_slide_provenance(tmp_path):
    soffice = shutil.which("soffice")
    deck = Presentation()
    for title, body in (
        (
            "Processes",
            "A process is a program in execution. Each process has its own address space. "
            "The operating system saves process state in a process control block during a context switch.",
        ),
        (
            "Scheduling",
            "Round robin scheduling gives each ready process a fixed time quantum. "
            "When the quantum expires, the running process goes to the back of the ready queue.",
        ),
    ):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    input_pptx = tmp_path / "fixture.pptx"
    deck.save(input_pptx)
    conversion_profile = tmp_path / "conversion-profile"
    subprocess.run(
        [
            soffice,
            f"-env:UserInstallation={conversion_profile.as_uri()}",
            "--headless",
            "--convert-to",
            "ppt",
            "--outdir",
            str(tmp_path),
            str(input_pptx),
        ],
        check=True,
        capture_output=True,
        text=True,
        timeout=60,
    )
    legacy_ppt = tmp_path / "fixture.ppt"
    assert legacy_ppt.is_file()

    data_dir = tmp_path / "data"
    settings = Settings(
        _env_file=None,
        pocketta_data_dir=data_dir,
        libreoffice_path=soffice,
        model_provider="lm_studio",
        lm_studio_model_id="",
        lm_studio_api_key="",
        openai_api_key="",
        embedding_api_key="",
        knowledge_infer_topics=False,
    )
    database = Database(settings.database_path)
    database.initialize()
    store = CourseStore(database)
    store.initialize()
    course = store.create_course("Legacy PPT test", "Disposable integration fixture")
    source_id = str(uuid.uuid4())
    source_dir = data_dir / "sources" / source_id
    source_dir.mkdir(parents=True)
    source_path = source_dir / "source.ppt"
    shutil.copyfile(legacy_ppt, source_path)
    with database._connect() as connection:
        connection.execute(
            "INSERT INTO sources(id,course_id,name,type,file_path,status,created_at) VALUES(?,?,?,?,?,?,?)",
            (
                source_id,
                course["id"],
                "Synthetic two-slide legacy PPT",
                "ppt",
                str(source_path),
                "queued",
                datetime.now(UTC).isoformat(),
            ),
        )

    processor = DocumentProcessor(
        settings, store, knowledge=KnowledgeService(settings, store)
    )
    await processor.process(source_id, generation_only=False)

    assert store.source(source_id)["status"] == "ready"
    units = store.units(source_id)
    assert [unit["location"] for unit in units] == ["slide:1", "slide:2"]
    assert "address space" in units[0]["text"].casefold()
    assert "time quantum" in units[1]["text"].casefold()
    assert source_path.with_suffix(".pptx").is_file()
    assert source_path.with_suffix(".pdf").is_file()
